import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('/opt/sapply-klm/attached_assets/КД_18_05_26_Липковская_1779190435513.xls', data_only=True)
ws = wb.active
print(f'Sheet: {ws.title}')
print(f'Rows: {ws.max_row}, Cols: {ws.max_column}')
print()
print('First 15 rows:')
for r in range(1, min(16, ws.max_row+1)):
    row = []
    for c in range(1, min(12, ws.max_column+1)):
        v = ws.cell(r, c).value
        row.append(str(v)[:30] if v else '')
    print(f'Row {r:2}: {row}')
